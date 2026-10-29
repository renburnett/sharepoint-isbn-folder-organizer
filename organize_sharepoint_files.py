from utils.logging_service import LogService
from utils.api_service import ApiService
from config import SID
import time
import ctypes
import os
import re
from pathlib import Path
import ebooklib
import pymupdf
import pytesseract
import logging

pytesseract.pytesseract.tesseract_cmd = (
    rf"C:\Users\{SID}\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"
)
from bs4 import BeautifulSoup
from docx import Document
from ebooklib import epub
from PIL import Image
from pypdf import PdfReader
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.edge.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    ElementClickInterceptedException,
    StaleElementReferenceException,
    TimeoutException,
)
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.action_chains import ActionChains


class TextbookCatalogingService:
    SHAREPOINT_URL = (
        "https://csulb.sharepoint.com/sites/DSA-BMAC/Shared%20Documents/"
        "Forms/AllItems.aspx?id=%2Fsites%2FDSA-BMAC%2FShared%20Documents"
        "%2F00%20Programs%20Workflow%2FMaster%20Documents%2FATP"
        "%2FAlternative%20Formats%2FTextbooks%2FArchives"
        "&sortField=BookTitle&isAscending=true"
    )
    SUPPORTED_EXTENSIONS = {"pdf", "tif", "tiff", "epub", "docx", "doc"}
    PHONE_PATTERN = re.compile(
        r"(?<!\d)\(?\d{3}\)?[^\dA-Za-z]+\d{3}[^\dA-Za-z]+\d{4}(?!\d)"
    )
    ISBN_13_PATTERN = re.compile(r"(?<!\d)97[89](?:[^\dA-Za-z]?\d){10}(?!\d)")
    ISBN_10_PATTERN = re.compile(r"(?<!\d)(?:\d[^\dA-Za-z]?){9}[\dXx](?!\d)")
    MINIMUM_ROWS_TO_LOAD = 12

    def get_listing_signature(self, driver):
        elements = driver.find_elements(
            By.CSS_SELECTOR,
            '[role="row"][data-automationid^="row-"] '
            '[data-automationid="field-LinkFilename"] '
            '[data-id="heroField"]',
        )

        names = []
        for element in elements:
            name = (element.get_attribute("title") or element.text).strip()
            if name:
                names.append(name)

        return tuple(names)

    def wait_for_folder_listing(self, driver, previous_signature=None, timeout=10):
        deadline = time.monotonic() + timeout
        last_signature = None
        stable_since = None

        empty_phrases = (
            "this folder is empty",
            "no items found",
            "no files found",
            "no items to show",
            "drag files here",
        )
        busy_selector = (
            '[aria-busy="true"], [role="progressbar"], '
            '[data-automationid*="loading"]'
        )

        while time.monotonic() < deadline:
            try:
                signature = self.get_listing_signature(driver)

                body_text = driver.find_element(By.TAG_NAME, "body").text.casefold()

                busy_elements = driver.find_elements(By.CSS_SELECTOR, busy_selector)
                busy = any(element.is_displayed() for element in busy_elements)
                empty = any(phrase in body_text for phrase in empty_phrases)

            except StaleElementReferenceException:
                # SharePoint refreshed the page while Selenium was reading it.
                time.sleep(0.4)
                continue

            has_results_or_empty_message = bool(signature) or empty
            changed_from_parent = (
                previous_signature is None
                or not previous_signature
                or signature != previous_signature
            )

            if has_results_or_empty_message and changed_from_parent and not busy:
                if signature == last_signature:
                    if (
                        stable_since is not None
                        and time.monotonic() - stable_since >= 2
                    ):
                        return signature
                else:
                    last_signature = signature
                    stable_since = time.monotonic()
            else:
                last_signature = None
                stable_since = None

            time.sleep(0.4)

        raise TimeoutError("The SharePoint folder listing did not finish loading.")

    def inspect_open_folder(self, driver, wait, current_folder):
        supported_extensions = set(self.SUPPORTED_EXTENSIONS) | {"doc"}
        format_priority = ("pdf", "tif", "tiff", "epub", "docx", "doc")

        folders = self.get_folders_to_update(driver)

        def normalized_title(item):
            return (
                str(item.get("title") or "")
                .casefold()
                .replace("_", " ")
                .replace("-", " ")
            )

        def title_tokens(item):
            return set(re.findall(r"[a-z0-9]+", normalized_title(item)))

        def formats_in(item):
            formats = set()
            format_text = str(item.get("format") or "").casefold()
            format_tokens = set(re.findall(r"[a-z0-9]+", format_text))
            formats.update(format_tokens & supported_extensions)

            title = str(item.get("title") or "")
            extension = os.path.splitext(title)[1].casefold().lstrip(".")
            if extension:
                formats.add(extension if extension in supported_extensions else "other")
            elif not item.get("is_folder") and format_tokens and not formats:
                formats.add("other")

            title_format_tokens = title_tokens(item)
            formats.update(title_format_tokens & supported_extensions)
            return formats

        def record_formats(items):
            formats = set(current_folder.get("format") or [])
            for item in items:
                formats.update(formats_in(item))
            current_folder["format"] = sorted(formats)

        def format_for(item):
            formats = formats_in(item)
            for extension in format_priority:
                if extension in formats:
                    return extension
            return "other" if "other" in formats else None

        def listing(items):
            return [
                f"{'[folder]' if item.get('is_folder') else '[file]'} "
                f"{item.get('title') or '(untitled)'}"
                for item in items
            ]

        def isbn_from_title(item, pattern):
            title = str(item.get("title") or "")
            file_format = format_for(item)
            item["format"] = item.get("format") or file_format

            for match in pattern.finditer(title):
                if self.PHONE_PATTERN.fullmatch(match.group()):
                    continue
                isbn = re.sub(r"[^0-9Xx]", "", match.group()).upper()
                return isbn, file_format

            return None, file_format

        # Record formats from the current folder's files and subfolders.
        record_formats(folders)

        format_folders = [
            item
            for item in folders
            if item.get("is_folder") and formats_in(item) & supported_extensions
        ]
        unexpected_folders = [
            item
            for item in folders
            if item.get("is_folder") and item not in format_folders
        ]

        contents = []
        used_unexpected_nesting = False

        if unexpected_folders and not format_folders:
            unexpected_folder = unexpected_folders[0]
            self.logger.write(
                "Unexpected subfolder; inspecting its listing:",
                unexpected_folder.get("title"),
                level=logging.INFO,
            )

            self.open_folder(driver, wait, unexpected_folder)
            unexpected_contents = self.get_folders_to_update(driver)
            record_formats(unexpected_contents)

            # Allow one extra level when the unexpected folder contains only one folder.
            child_folders = [
                item for item in unexpected_contents if item.get("is_folder")
            ]

            if len(unexpected_contents) == 1 and len(child_folders) == 1:
                child_folder = child_folders[0]
                self.logger.write(
                    "Opening its single subfolder:",
                    child_folder.get("title"),
                    level=logging.INFO,
                )

                self.open_folder(driver, wait, child_folder)
                contents = self.get_folders_to_update(driver)
                record_formats(contents)

                if any(item.get("is_folder") for item in contents):
                    self.logger.write(
                        f"Unexpected deeper folder(s) inside "
                        f"{child_folder['title']!r}. "
                        f"Contents: {listing(contents)}",
                        level=logging.WARNING,
                    )
                    # TODO: recurse through all folder levels at some point?
                    return None

                used_unexpected_nesting = True

            elif child_folders:
                self.logger.write(
                    f"Unexpected subfolder structure in "
                    f"{unexpected_folder.get('title')!r}. "
                    f"Contents: {listing(unexpected_contents)}",
                    level=logging.WARNING,
                )
                # TODO: recurse through all folder levels at some point?
                return None

            else:
                # This folder contains files directly, so inspect those files.
                contents = unexpected_contents
                record_formats(contents)
                used_unexpected_nesting = True

        def folder_priority(item):
            item_formats = formats_in(item)
            first_format = next(
                (ext for ext in format_priority if ext in item_formats),
                "other",
            )
            format_rank = (
                format_priority.index(first_format)
                if first_format in format_priority
                else 99
            )
            is_split = "split" in normalized_title(item)
            split_pdf = is_split and "pdf" in item_formats
            return (
                (
                    0
                    if split_pdf
                    else 1 if first_format == "pdf" else 2 if is_split else 3
                ),
                format_rank,
            )

        supported_format_folder = None

        if not used_unexpected_nesting:
            supported_format_folder = (
                min(format_folders, key=folder_priority) if format_folders else None
            )

            if supported_format_folder:
                self.logger.write(
                    "Opening format folder:", supported_format_folder["title"]
                )
                self.open_folder(driver, wait, supported_format_folder)
                contents = self.get_folders_to_update(driver)
                record_formats(contents)

                if any(item.get("is_folder") for item in contents):
                    self.logger.write(
                        f"Unexpected deeper folder(s) inside "
                        f"{supported_format_folder['title']!r}. "
                        f"Contents: {listing(contents)}",
                        level=logging.WARNING,
                    )
                    # TODO: recurse through all folder levels once recursive handling is added.
                    return None

        # Check PDFs first, then other supported formats; ISBN-13 precedes ISBN-10.
        search_items = sorted(
            folders + contents,
            key=lambda item: (
                0
                if format_for(item) == "pdf"
                else 1 if format_for(item) in supported_extensions else 2
            ),
        )

        for pattern in (self.ISBN_13_PATTERN, self.ISBN_10_PATTERN):
            for item in search_items:
                isbn, file_format = isbn_from_title(item, pattern)
                if isbn:
                    self.logger.write(
                        f"Found ISBN in {file_format or 'untyped'} item "
                        f"{item.get('title')!r}: {isbn}"
                    )
                    catalog = self.api.lookup_isbn(isbn)
                    if catalog:
                        return catalog

        if supported_format_folder or used_unexpected_nesting:
            current_items = contents
        else:
            current_items = folders
            self.logger.write(
                "No supported-format folder was found; checking the current folder.",
                level=logging.WARNING,
            )

        matter_files = [
            item
            for item in current_items
            if not item.get("is_folder")
            and "matter" in normalized_title(item)
            and format_for(item) in {"pdf", "tif", "tiff", "docx", "epub"}
        ]
        matter_files.sort(
            key=lambda item: (
                0
                if format_for(item) == "pdf"
                else 1 if format_for(item) in supported_extensions else 2
            )
        )

        front_matter_file = next(
            (item for item in matter_files if "front" in title_tokens(item)),
            None,
        )
        end_matter_file = next(
            (item for item in matter_files if {"end", "back"} & title_tokens(item)),
            None,
        )
        matter_file = (
            front_matter_file or end_matter_file or next(iter(matter_files), None)
        )

        if not matter_file:
            self.logger.write(
                f"No front-matter or end-matter file was found. "
                f"Items: {listing(current_items)}",
                level=logging.WARNING,
            )
            return None

        self.logger.write("Found a possible matter file:", matter_file["title"])
        downloaded_path = self.download_sharepoint_file(driver, wait, matter_file)
        self.logger.write("Downloaded:", downloaded_path)

        matter_tokens = title_tokens(matter_file)
        if "front" in matter_tokens:
            matter_type = "front"
        elif {"end", "back"} & matter_tokens:
            matter_type = "end"
        else:
            matter_type = "front"

        isbn = self.find_isbn_in_downloaded_file(downloaded_path, matter_type)

        if isbn:
            self.logger.write(
                f"Found ISBN in downloaded file {matter_file['title']!r}: {isbn}"
            )
            return self.api.lookup_isbn(isbn)

        self.logger.write(
            f"No ISBN found in the downloaded file {matter_file['title']!r}.",
            level=logging.WARNING,
        )
        return None

    def open_folder(self, driver, wait, folder):
        old_row = folder["_elements"]["row"]
        old_url = driver.current_url
        old_signature = self.get_listing_signature(driver)
        wait.until(EC.element_to_be_clickable(folder["_elements"]["title"])).click()
        wait.until(EC.any_of(EC.staleness_of(old_row), EC.url_changes(old_url)))
        self.wait_for_folder_listing(driver, previous_signature=old_signature)

    def open_first_eligible_folder(self, driver, wait, open_subfolders=True):
        self.wait_for_folder_listing(driver)
        wait.until(
            lambda d: len(
                d.find_elements(
                    By.CSS_SELECTOR, '[role="row"][data-automationid^="row-"]'
                )
            )
            > self.MINIMUM_ROWS_TO_LOAD
        )
        items = self.get_folders_to_update(driver)
        excluded_names = ("abueg", "izaiah", "burnett", "-sa")

        for item in items:
            modified_by = (item.get("modified_by") or "").casefold()
            formats = (item.get("format") or "").strip()
            book_title = (item.get("book_title") or "").strip()
            excluded = any(name in modified_by for name in excluded_names)
            if (
                item.get("is_folder")
                and not excluded
                and not formats
                and not book_title
            ):
                if open_subfolders:
                    self.open_folder(driver, wait, item)
                return item
        print("No folder matches those conditions.")
        self.logger.write(
            f"No folder matches the conditions for opening. Folders: {items}",
            level=logging.WARNING,
        )
        return None

    def click_safely(self, driver, wait, locator):
        try:
            wait.until(EC.element_to_be_clickable(locator)).click()
        except ElementClickInterceptedException:
            # Clear the profile card that can appear on hover.
            ActionChains(driver).send_keys(Keys.ESCAPE).perform()
            wait.until(
                EC.invisibility_of_element_located(
                    (
                        By.CSS_SELECTOR,
                        'lpc-card[role="dialog"][aria-label="Profile Card"]',
                    )
                )
            )
            wait.until(EC.element_to_be_clickable(locator)).click()

    def read_cell(self, driver, row, selector, use_title_attribute=False):
        matches = row.find_elements(By.CSS_SELECTOR, selector)
        if not matches:
            return None, None
        element = matches[0]
        text = element.get_attribute("title") if use_title_attribute else None
        if not text:
            text = driver.execute_script(
                "return (arguments[0].innerText || "
                "arguments[0].textContent || '').trim();",
                element,
            )
        return text or None, element

    def get_folders_to_update(self, driver):
        for attempt in range(3):
            try:
                return self._read_folders_to_update_once(driver)
            except StaleElementReferenceException:
                if attempt == 2:
                    raise
                time.sleep(0.5)

    def _read_folders_to_update_once(self, driver):
        rows = driver.find_elements(
            By.CSS_SELECTOR, '[role="row"][data-automationid^="row-"]'
        )
        print("DOM rows:", len(rows))
        items = []
        for row in rows:
            title, title_element = self.read_cell(
                driver,
                row,
                '[data-automationid="field-LinkFilename"] ' '[data-id="heroField"]',
                use_title_attribute=True,
            )
            if not title_element:
                continue
            modified_by, modified_by_cell = self.read_cell(
                driver, row, '[data-automationid="field-Editor"]'
            )
            formats, formats_cell = self.read_cell(
                driver, row, '[data-automationid="field-Formats"]'
            )
            book_title, book_title_cell = self.read_cell(
                driver, row, '[data-automationid="field-BookTitle"]'
            )
            author, author_cell = self.read_cell(
                driver, row, '[data-automationid="field-Author0"]'
            )
            publisher, publisher_cell = self.read_cell(
                driver, row, '[data-automationid="field-Publisher"]'
            )
            edition, edition_cell = self.read_cell(
                driver, row, '[data-automationid="field-Edition"]'
            )
            is_folder = bool(
                row.find_elements(
                    By.CSS_SELECTOR,
                    '[data-automationid="field-DocIcon"] '
                    '[aria-label="Yellow folder"]',
                )
            )
            items.append(
                {
                    "title": title,
                    "modified_by": modified_by,
                    "format": formats,
                    "book_title": book_title,
                    "author": author,
                    "publisher": publisher,
                    "edition": edition,
                    "is_folder": is_folder,
                    "_elements": {
                        "row": row,
                        "title": title_element,
                        "modified_by": modified_by_cell,
                        "format": formats_cell,
                        "book_title": book_title_cell,
                        "author": author_cell,
                        "publisher": publisher_cell,
                        "edition": edition_cell,
                    },
                }
            )
        print(f"Found {len(items)} item(s).")
        return items

    def get_folder_by_title(self, driver, title):
        return next(
            (
                item
                for item in self.get_folders_to_update(driver)
                if item.get("is_folder") and item.get("title") == title
            ),
            None,
        )

    def rename_sharepoint_folder(self, driver, wait, folder, new_title):
        old_title = folder["title"]
        original_url = driver.current_url
        original_signature = self.get_listing_signature(driver)
        row_ref = [folder["_elements"]["row"]]

        editor_selectors = (
            'input:not([type="checkbox"])',
            "textarea",
            '[contenteditable="true"]',
            '[role="textbox"]',
        )

        def find_visible_editor(driver):
            containers = [row_ref[0]]

            for locator in (
                (By.ID, "CommonEditorCalloutId"),
                (By.CSS_SELECTOR, '[role="dialog"]'),
            ):
                containers.extend(
                    element
                    for element in driver.find_elements(*locator)
                    if element.is_displayed()
                )

            for container in containers:
                try:
                    for selector in editor_selectors:
                        for element in container.find_elements(
                            By.CSS_SELECTOR, selector
                        ):
                            if element.is_displayed():
                                return element
                except StaleElementReferenceException:
                    continue

            return False

        title_cell = row_ref[0].find_element(
            By.CSS_SELECTOR,
            '[data-automationid="field-LinkFilename"]',
        )
        title_cell.click()

        try:
            # Give the first click a moment to open the editor.
            editor = WebDriverWait(driver, 2).until(find_visible_editor)
        except TimeoutException:
            # Don't click again if the first click navigated or changed the listing.
            if (
                driver.current_url != original_url
                or self.get_listing_signature(driver) != original_signature
            ):
                raise TimeoutException(
                    "Clicking the title changed the listing instead of opening its editor."
                )

            # Refresh the row reference, then try the title cell once more.
            current_folder = self.get_folder_by_title(driver, old_title)
            if not current_folder:
                raise TimeoutException(
                    f"Could not find the folder {old_title!r} after selecting it."
                )

            folder.update(current_folder)
            row_ref[0] = folder["_elements"]["row"]

            title_cell = row_ref[0].find_element(
                By.CSS_SELECTOR,
                '[data-automationid="field-LinkFilename"]',
            )
            title_cell.click()
            editor = wait.until(find_visible_editor)

        editor.click()
        editor.send_keys(Keys.CONTROL, "a")
        editor.send_keys(new_title)
        editor.send_keys(Keys.ENTER)

        wait.until(
            EC.any_of(
                EC.staleness_of(editor),
                EC.invisibility_of_element(editor),
            )
        )
        wait.until(lambda d: new_title in self.get_listing_signature(d))

        folder["title"] = new_title

        updated_folder = self.get_folder_by_title(driver, new_title)
        if updated_folder:
            folder.update(updated_folder)

    def update_folder_formats(self, driver, wait, folder, formats):
        format_names = {
            "pdf": "PDF",
            "epub": "Epub",
            "docx": "Docx",
            "tif": "TIF",
            "tiff": "TIF",
        }
        selected_formats = list(
            dict.fromkeys(
                format_names[extension.casefold()]
                for extension in formats
                if extension.casefold() in format_names
            )
        )
        if not selected_formats:
            return

        folder = self.get_folder_by_title(driver, folder["title"]) or folder
        folder_title = folder["title"]
        original_url = driver.current_url
        original_signature = self.get_listing_signature(driver)
        format_cell = folder["_elements"]["format"]

        try:
            wait.until(EC.element_to_be_clickable(format_cell)).click()
            WebDriverWait(driver, 2).until(
                EC.visibility_of_element_located((By.ID, "CommonEditorCalloutId"))
            )
        except TimeoutException:
            if (
                driver.current_url != original_url
                or self.get_listing_signature(driver) != original_signature
            ):
                raise TimeoutException(
                    "Clicking the Formats cell changed the listing instead "
                    "of opening its editor."
                )

            refreshed_folder = self.get_folder_by_title(driver, folder_title)
            if not refreshed_folder:
                raise TimeoutException(
                    f"Could not find folder {folder_title!r} after selecting "
                    "the Formats cell."
                )

            folder.update(refreshed_folder)
            format_cell = folder["_elements"]["format"]
            self.logger.write(
                "Formats cell was selected; clicking it again to open the editor."
            )
            wait.until(EC.element_to_be_clickable(format_cell)).click()
            wait.until(
                EC.visibility_of_element_located((By.ID, "CommonEditorCalloutId"))
            )

        for format_name in selected_formats:
            callout = wait.until(
                EC.visibility_of_element_located((By.ID, "CommonEditorCalloutId"))
            )
            option = callout.find_element(
                By.CSS_SELECTOR, f'[role="option"][title="{format_name}"]'
            )
            driver.execute_script("arguments[0].click();", option)

        callout = wait.until(
            EC.visibility_of_element_located((By.ID, "CommonEditorCalloutId"))
        )
        save_button = callout.find_element(By.ID, "acceptButtonInCommonEditorCallout")
        driver.execute_script("arguments[0].click();", save_button)
        wait.until(EC.invisibility_of_element_located((By.ID, "CommonEditorCalloutId")))

    def update_folder_text_field(self, driver, wait, folder, field_name, value):
        if not value:
            return

        folder = self.get_folder_by_title(driver, folder["title"]) or folder
        folder_title = folder["title"]
        original_url = driver.current_url
        original_signature = self.get_listing_signature(driver)
        row_ref = [folder["_elements"]["row"]]

        editor_selectors = (
            'input:not([type="checkbox"])',
            "textarea",
            '[contenteditable="true"]',
            '[role="textbox"]',
        )

        def find_visible_editor(driver):
            containers = []

            for locator in (
                (By.ID, "CommonEditorCalloutId"),
                (By.CSS_SELECTOR, '[role="dialog"]'),
            ):
                containers.extend(
                    element
                    for element in driver.find_elements(*locator)
                    if element.is_displayed()
                )

            try:
                if row_ref[0].is_displayed():
                    containers.append(row_ref[0])
            except StaleElementReferenceException:
                refreshed_folder = self.get_folder_by_title(driver, folder_title)
                if refreshed_folder:
                    folder.update(refreshed_folder)
                    row_ref[0] = folder["_elements"]["row"]
                    containers.append(row_ref[0])

            for container in containers:
                try:
                    for selector in editor_selectors:
                        for candidate in container.find_elements(
                            By.CSS_SELECTOR, selector
                        ):
                            if candidate.is_displayed():
                                return candidate
                except StaleElementReferenceException:
                    continue

            return False

        cell = folder["_elements"][field_name]
        wait.until(EC.element_to_be_clickable(cell)).click()

        try:
            # Give the first click a moment to open the editor.
            editor = WebDriverWait(driver, 2).until(find_visible_editor)
        except TimeoutException:
            # Retry only if the first click left the listing unchanged.
            if (
                driver.current_url != original_url
                or self.get_listing_signature(driver) != original_signature
            ):
                raise TimeoutException(
                    f"Clicking the {field_name} cell changed the listing "
                    "instead of opening its editor."
                )

            refreshed_folder = self.get_folder_by_title(driver, folder_title)
            if not refreshed_folder:
                raise TimeoutException(
                    f"Could not find folder {folder_title!r} after selecting "
                    f"the {field_name} cell."
                )

            folder.update(refreshed_folder)
            row_ref[0] = folder["_elements"]["row"]
            cell = folder["_elements"][field_name]

            wait.until(EC.element_to_be_clickable(cell)).click()
            editor = wait.until(find_visible_editor)

        editor.click()
        editor.send_keys(Keys.CONTROL, "a")
        editor.send_keys(str(value))
        editor.send_keys(Keys.ENTER)

        wait.until(
            EC.any_of(
                EC.staleness_of(editor),
                EC.invisibility_of_element(editor),
            )
        )

    def update_folder_from_catalog(self, driver, wait, folder, catalog):
        isbn_value = catalog.get("isbn_13") or catalog.get("isbn_10")

        if isinstance(isbn_value, (list, tuple)):
            isbn_value = next((value for value in isbn_value if value), None)

        isbn = re.sub(r"[^0-9Xx]", "", str(isbn_value)).upper() if isbn_value else None

        title = (
            catalog.get("title")
            or catalog.get("full_title")
            or catalog.get("short_title")
        )

        print("isbn", isbn)
        print("title", title)

        if isbn and isbn != folder["title"]:
            self.rename_sharepoint_folder(driver, wait, folder, isbn)

        input("\nPress Enter to continue...")

        # TODO: figure out why formats are not updating
        self.update_folder_formats(driver, wait, folder, folder.get("format") or [])

        authors = (
            catalog.get("by_statement")
            or catalog.get("authors")
            or catalog.get("contributions")
        )

        if isinstance(authors, list):
            authors = "; ".join(
                author.get("name", "") if isinstance(author, dict) else str(author)
                for author in authors
            )

        publishers = catalog.get("publishers") or catalog.get("publisher")
        if isinstance(publishers, list):
            publishers = "; ".join(map(str, publishers))

        self.update_folder_text_field(driver, wait, folder, "book_title", title)
        self.update_folder_text_field(driver, wait, folder, "author", authors)
        self.update_folder_text_field(driver, wait, folder, "publisher", publishers)
        self.update_folder_text_field(
            driver, wait, folder, "edition", catalog.get("edition_name")
        )

    def download_sharepoint_file(self, driver, wait, item):
        before = set(self.download_dir.iterdir())
        title_element = item["_elements"]["title"]
        wait.until(EC.visibility_of(title_element))
        driver.execute_script(
            "arguments[0].scrollIntoView({block: 'center'});", title_element
        )
        ActionChains(driver).move_to_element(title_element).pause(0.3).perform()
        row = title_element.find_element(By.XPATH, "./ancestor::*[@role='row'][1]")
        more_actions = row.find_element(
            By.CSS_SELECTOR, 'button[data-automationid="moreActionsHeroField"]'
        )
        print("Clicking More Actions for:", item["title"])
        print("more_actions", more_actions)
        wait.until(EC.visibility_of(more_actions))
        driver.execute_script("arguments[0].click();", more_actions)
        download_button_locator = (
            By.CSS_SELECTOR,
            'button[data-automationid="downloadCommand"]',
        )
        download_button = wait.until(
            EC.visibility_of_element_located(download_button_locator)
        )
        driver.execute_script("arguments[0].click();", download_button)
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline:
            new_files = [
                path
                for path in self.download_dir.iterdir()
                if path not in before and path.is_file()
            ]
            partial_download = any(
                path.name.endswith((".crdownload", ".tmp")) for path in new_files
            )
            completed_files = [
                path
                for path in new_files
                if not path.name.endswith((".crdownload", ".tmp"))
            ]
            if completed_files and not partial_download:
                return max(completed_files, key=lambda path: path.stat().st_mtime)
            time.sleep(0.5)
        raise TimeoutError("The SharePoint file did not finish downloading.")

    def find_isbn_in_downloaded_file(self, file_path, matter_type):
        extension = file_path.suffix.casefold()
        text_parts = []
        if extension == ".pdf":
            reader = PdfReader(str(file_path))
            pdf_for_ocr = pymupdf.open(str(file_path))
            if matter_type == "end":
                start = max(0, len(reader.pages) - 8)
                page_numbers = range(start, len(reader.pages))
            else:
                page_numbers = range(min(8, len(reader.pages)))
            for page_number in page_numbers:
                text = reader.pages[page_number].extract_text() or ""
                # OCR the page if it has no extractable text.
                if not text.strip():
                    pixmap = pdf_for_ocr[page_number].get_pixmap(dpi=200)
                    image = Image.frombytes(
                        "RGB", (pixmap.width, pixmap.height), pixmap.samples
                    )
                    text = pytesseract.image_to_string(image)
                text_parts.append(text)
        elif extension in {".tif", ".tiff"}:
            with Image.open(file_path) as image:
                frame_count = getattr(image, "n_frames", 1)
                if matter_type == "end":
                    start = max(0, frame_count - 8)
                    frame_numbers = range(start, frame_count)
                else:
                    frame_numbers = range(min(8, frame_count))
                for frame_number in frame_numbers:
                    image.seek(frame_number)
                    text_parts.append(pytesseract.image_to_string(image.copy()))
        elif extension == ".docx":
            document = Document(str(file_path))
            text_parts.extend(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    text_parts.append(" ".join(cell.text for cell in row.cells))
        elif extension == ".epub":
            book = epub.read_epub(str(file_path))
            for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                soup = BeautifulSoup(item.get_content(), "html.parser")
                text_parts.append(soup.get_text(" ", strip=True))
        else:
            raise ValueError(f"Unsupported file type: {extension}")
        text = "\n".join(text_parts)

        # Search for ISBN-13 across the text before trying ISBN-10.
        for pattern in (self.ISBN_13_PATTERN, self.ISBN_10_PATTERN):
            for match in pattern.finditer(text):
                if self.PHONE_PATTERN.fullmatch(match.group()):
                    continue
                return re.sub(r"[^0-9Xx]", "", match.group()).upper()

        return None

    def __init__(self):
        self.logger = LogService(
            Path(__file__).with_name("organize_sharepoint_files.log")
        )
        self.api = ApiService(log_service=self.logger)
        driver = None

        try:
            # Use the Windows Downloads folder.
            self.download_dir = Path.home() / "Downloads"
            self.download_dir.mkdir(exist_ok=True)

            options = Options()
            options.add_experimental_option(
                "prefs",
                {
                    "download.default_directory": str(self.download_dir.resolve()),
                    "download.prompt_for_download": False,
                },
            )

            driver = webdriver.Edge(options=options)
            wait = WebDriverWait(driver, 10)

            screen_width = ctypes.windll.user32.GetSystemMetrics(0)
            screen_height = ctypes.windll.user32.GetSystemMetrics(1)

            driver.get(self.SHAREPOINT_URL)
            driver.set_window_rect(
                x=0,
                y=0,
                width=screen_width,
                height=screen_height,
            )
            self.logger.write("\nEdge opened your SharePoint folder.")

            # Load all folders.
            see_more = (
                By.CSS_SELECTOR,
                'div[role="gridcell"][class*="seeMoreCell"] button',
            )
            self.click_safely(driver, wait, see_more)
            first_folder = self.open_first_eligible_folder(driver, wait)

            if first_folder:
                book_folder = first_folder
                result = self.inspect_open_folder(driver, wait, book_folder)

                self.logger.write("Catalog lookup result:", result)

                driver.get(self.SHAREPOINT_URL)
                self.click_safely(driver, wait, see_more)

                detected_formats = book_folder.get("format") or []
                first_folder = self.open_first_eligible_folder(driver, wait, False)

                if (
                    first_folder
                    and result
                    and str(result.get("key", "")).startswith("/books/")
                ):
                    first_folder["format"] = detected_formats
                    self.update_folder_from_catalog(driver, wait, first_folder, result)

            input("\nPress Enter to close Edge...")

        except (TimeoutException, TimeoutError) as exc:
            self.logger.write(
                f"\nTimed out waiting for SharePoint: {exc}",
                level=logging.ERROR,
            )
            self.logger.write("Closing Edge and stopping the script.")
        except Exception as exc:
            self.logger.exception(f"Unexpected error ({type(exc).__name__}): {exc}")
        finally:
            if driver:
                driver.quit()


if __name__ == "__main__":
    TextbookCatalogingService()
